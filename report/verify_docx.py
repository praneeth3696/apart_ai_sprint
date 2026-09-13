"""
verify_docx.py

Checks the built DOCX against `analysis/e1_stats.json` and against the
submission requirements. Run it after `build_docx.py`, and again before
submitting.

It exists for the same reason `check_numbers.py` does: the markdown sections
and the DOCX are two renderings of one dataset, and a number can go stale in
either. This one reads the actual .docx, so it catches a build that silently
used an older stats file.
"""
from __future__ import annotations

import json
import pathlib
import re
import sys
import zipfile

import docx

REPO = pathlib.Path(__file__).resolve().parent.parent
DOCX = REPO / "report" / "HF-Replay-Recon_AI-Incident-Response-Sprint.docx"
STATS = json.loads((REPO / "analysis" / "e1_stats.json").read_text(encoding="utf-8"))
PER, E3 = STATS["per_model"], STATS["e3"]["per_model"]
WORD_RE = r"[A-Za-z0-9][A-Za-z0-9'’./%$-]*"

fails: list[str] = []
oks = 0


def check(cond: bool, msg: str) -> None:
    global oks
    if cond:
        oks += 1
    else:
        fails.append(msg)


def short(m: str) -> str:
    n = m.split(":", 1)[-1].split("/")[-1]
    return n[:-7] if n.endswith("-latest") else n


def main() -> int:
    global oks
    if not DOCX.is_file():
        print(f"no {DOCX} — run report/build_docx.py first", file=sys.stderr)
        return 2

    check(zipfile.ZipFile(DOCX).testzip() is None, "docx is a corrupt zip")
    d = docx.Document(str(DOCX))
    text = " ".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                text += " " + c.text

    # ---- submission requirements ------------------------------------------
    title_cell = d.tables[0].rows[0].cells[0]
    abstract_cell = d.tables[0].rows[1].cells[0]
    ps = [p.text for p in abstract_cell.paragraphs if p.text.strip()]
    n_abs = len(re.findall(WORD_RE, " ".join(ps[1:])))
    check(n_abs <= 150, f"abstract is {n_abs} words, limit 150")
    check(any(p.style.name == "Title" and p.text.strip()
              for p in title_cell.paragraphs), "no Title-styled title")
    check("@" in title_cell.text, "no author emails in the title block")
    heads = [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Heading")]
    for required in ("1. Introduction", "2. Related Work", "3. Methods", "4. Results",
                     "5. Discussion and Limitations", "Limitations", "Future Work",
                     "6. Conclusion", "Code and Data", "References",
                     "LLM Usage Statement"):
        check(any(h == required or h.startswith(required) for h in heads),
              f"template section missing: {required}")
    check(any("Dual-Use" in h for h in heads),
          "the required Limitations and Dual-Use appendix is missing")
    check(len(d.inline_shapes) >= 2, "fewer than two figures")
    check(not re.search(r"\[(First|Second|Third) contribution|\[Reference \d|\[Link to",
                        text), "template placeholder text left in the document")
    check("**" not in text, "unrendered markdown bold in the document")

    # ---- every model row against the data ---------------------------------
    e1_tbl = next(t for t in d.tables if len(t.columns) == 5 and "Fisher" in t.rows[0].cells[-1].text)
    seen = set()
    for r in e1_tbl.rows[1:]:
        name = r.cells[0].text.strip()
        seen.add(name)
        m = next((k for k in PER if short(k) == name), None)
        check(m is not None, f"E1 row '{name}' is not in the data")
        if not m:
            continue
        s = PER[m]
        h, b = s["milestone_hit_rate"], s["benign_false_page_rate"]
        check(f"{h['k']}/{h['n']}" in r.cells[2].text,
              f"{name}: milestone cell '{r.cells[2].text}' != {h['k']}/{h['n']}")
        if b["n"]:
            check(f"{b['k']}/{b['n']}" in r.cells[3].text,
                  f"{name}: benign cell '{r.cells[3].text}' != {b['k']}/{b['n']}")
    check(not (set(short(k) for k in PER) - seen),
          f"models absent from the E1 table: {sorted(set(short(k) for k in PER) - seen)}")

    e3_tbl = next(t for t in d.tables if len(t.columns) == 6 and "McNemar" in t.rows[0].cells[-1].text)
    for r in e3_tbl.rows[1:]:
        name = r.cells[0].text.strip()
        m = next((k for k in E3 if short(k) == name), None)
        check(m is not None, f"E3 row '{name}' is not in the data")
        if not m:
            continue
        s = E3[m]
        check(f"{s['incoherence_rate']['k']}/{s['incoherence_rate']['n']}" in r.cells[2].text,
              f"{name}: incoherence mismatch")
        check(f"{s['escalated']['k']}/{s['escalated']['n']}" in r.cells[3].text,
              f"{name}: escalated mismatch")

    # ---- pooled claims in the prose ---------------------------------------
    agg = lambda key, sub: sum(s[key][sub] for s in PER.values())
    e3agg = lambda key, sub: sum(s[key][sub] for s in E3.values())
    sig = [m for m, s in PER.items()
           if s["fisher_milestone_vs_benign"] is not None
           and s["fisher_milestone_vs_benign"] < 0.05]
    for label, needle in (
        ("pooled milestone", f"{agg('milestone_hit_rate','k')}/{agg('milestone_hit_rate','n')}"),
        ("pooled benign", f"{agg('benign_false_page_rate','k')}/{agg('benign_false_page_rate','n')}"),
        ("E3 incoherent", f"{e3agg('incoherence_rate','k')} of {e3agg('incoherence_rate','n')}"),
        ("E3 escalated", f"{e3agg('escalated','k')}/{e3agg('refused_to_assist','n')}"),
        ("significance count", f"{len(sig)} of {len(PER)}"),
    ):
        check(needle in text, f"{label}: '{needle}' not found in the document")

    print(f"{DOCX.name}: {oks} checks passed")
    if fails:
        for f in fails:
            print(f"  FAIL  {f}")
        return 1
    print("  all pass — the document matches the data and the template requirements")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
